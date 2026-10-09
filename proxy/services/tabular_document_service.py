"""Literal, bounded table snapshots. No row classification or domain inference."""
from __future__ import annotations

import csv
from datetime import date, datetime, time
import hashlib
import io
import json
import math
from pathlib import Path
import zipfile

TABLE_SUFFIXES = {'.xlsx', '.xlsm', '.csv'}
MAX_ROWS = 20_000
MAX_CELLS = 300_000
MAX_COLUMNS = 512
MAX_EXPANDED_BYTES = 64 * 1024 * 1024


def encode(value):
    return json.dumps(value, ensure_ascii=False, separators=(',', ':'), allow_nan=False, sort_keys=True)


def cell_value(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return {'kind': 'boolean', 'value': value}
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError('Таблица содержит нечисловое значение NaN/Infinity')
        return {'kind': 'number', 'value': str(value)}
    if isinstance(value, (datetime, date, time)):
        return {'kind': 'date', 'value': value.isoformat()}
    return {'kind': 'text', 'value': str(value)}


def read_table(path: Path, name: str | None = None) -> dict:
    """Capture every nonempty row, preserving coordinates, zeros and formulas.

    Empty cells are represented by absent coordinates, never by shifting cells.
    Bounds reject the whole source; a partial extraction cannot look complete.
    """
    path = Path(path)
    if path.suffix.lower() not in TABLE_SUFFIXES:
        raise ValueError('Поддерживаются XLSX, XLSM и CSV')
    if path.stat().st_size > MAX_EXPANDED_BYTES:
        raise ValueError('Таблица превышает предел чтения 64 МиБ')
    raw = path.read_bytes()
    if len(raw) > MAX_EXPANDED_BYTES:
        raise ValueError('Таблица превышает предел чтения 64 МиБ')
    rows, sheets = [], []
    count = 0

    def add_row(sheet_index, sheet_name, number, cells):
        nonlocal count
        if not cells:
            return
        count += len(cells)
        if len(rows) >= MAX_ROWS or count > MAX_CELLS:
            raise ValueError('Таблица превышает предел: 20000 строк или 300000 ячеек')
        rows.append({'id': f's{sheet_index}:r{number}', 'sheet': sheet_name,
                     'row': number, 'locator': f'{sheet_name}!R{number}', 'cells': cells})

    if path.suffix.lower() == '.csv':
        from backend.text_decoding import read_document_text
        from openpyxl.utils import get_column_letter
        text = read_document_text(path)
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        width = 0
        for number, values in enumerate(csv.reader(io.StringIO(text, newline=''), dialect), 1):
            if len(values) > MAX_COLUMNS or number > MAX_ROWS:
                raise ValueError('CSV превышает предел строк или столбцов')
            width = max(width, len(values))
            cells = [{'coordinate': f'{get_column_letter(col)}{number}', **cell_value(value)}
                     for col, value in enumerate(values, 1) if value != '']
            add_row(1, 'CSV', number, cells)
        sheets.append({'name': 'CSV', 'columns': width, 'merged_ranges': [], 'hidden': False})
    else:
        import openpyxl
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if len(archive.infolist()) > 10000 or sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise ValueError('Распакованная таблица превышает 64 МиБ')
        book = openpyxl.load_workbook(io.BytesIO(raw), data_only=False, keep_links=False)
        cached = None
        try:
            cached = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, keep_links=False)
            for index, sheet in enumerate(book, 1):
                if (sheet.max_row > MAX_ROWS or sheet.max_column > MAX_COLUMNS
                        or sheet.max_row * sheet.max_column > MAX_CELLS):
                    raise ValueError('Лист превышает предел строк или столбцов')
                sheets.append({'name': sheet.title, 'columns': sheet.max_column,
                               'merged_ranges': [str(item) for item in sheet.merged_cells.ranges],
                               'hidden': sheet.sheet_state != 'visible'})
                for row in sheet.iter_rows():
                    cells = []
                    for cell in row:
                        if cell.value is None or cell.value == '':
                            continue
                        if cell.data_type == 'f':
                            cached_cell = cached[sheet.title][cell.coordinate]
                            value = cached_cell.value
                            payload = {'kind': 'formula', 'formula': str(cell.value),
                                       'cached': ({'kind': 'error', 'value': str(value)}
                                                  if cached_cell.data_type == 'e' else cell_value(value)),
                                       'needs_recalculation': value is None or cached_cell.data_type == 'e'}
                        elif cell.data_type == 'e':
                            payload = {'kind': 'error', 'value': str(cell.value)}
                        else:
                            payload = cell_value(cell.value)
                        cells.append({'coordinate': cell.coordinate, **payload})
                    add_row(index, sheet.title, row[0].row, cells)
        finally:
            book.close()
            if cached is not None:
                cached.close()
    digest = hashlib.sha256(raw).hexdigest()
    # CSV decoding reads by path; reject replacement during extraction as well.
    if hashlib.sha256(path.read_bytes()).hexdigest() != digest:
        raise ValueError('Файл изменился во время чтения. Прикрепите его заново')
    return {'schema': 'les.table.v1', 'name': name or path.name, 'sha256': digest,
            'sheets': sheets, 'rows': rows, 'row_count': len(rows),
            'cell_policy': 'Coordinates are literal. Absent coordinates are blank. '
                           'Numbers are decimal strings; formula caches may be absent. '
                           'All nonempty rows, including headings, are included.'}


def preview(document: dict, max_chars: int) -> tuple[str, bool]:
    """A whole-row preview only; original values remain in the task snapshot."""
    parts = [f"Файл: {document['name']}", document['cell_policy'],
             f"Всего непустых строк: {document['row_count']}"]
    used = len('\n'.join(parts))
    shown = 0
    for row in document['rows']:
        line = f"{row['locator']}: " + encode(row['cells'])
        if used + len(line) + 160 > max_chars:
            break
        parts.append(line)
        used += len(line) + 1
        shown += 1
    truncated = shown < document['row_count']
    if truncated:
        parts.append(f"Показано {shown} из {document['row_count']} строк. Полное чтение — инструментом table_document.")
    return '\n'.join(parts), truncated
