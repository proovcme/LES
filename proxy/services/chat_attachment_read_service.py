"""Read attachments with row provenance and preserved server-owned originals."""
import asyncio
import logging
import os
from pathlib import Path
from typing import Any
from fastapi import HTTPException
logger = logging.getLogger(__name__)
_READ_ATTACH_MAX_CHARS = int(os.getenv("RAG_ATTACH_READ_MAX_CHARS", "18000"))

def _compact_cell(value: Any, *, max_len: int = 240) -> str:
    text = " ".join(str(value or "").replace("\u00a0", " ").split())
    if len(text) > max_len:
        return text[: max_len - 1].rstrip() + "…"
    return text

def _format_tabular_attachment_context(path: Path, original_name: str, *, max_chars: int) -> tuple[str, bool] | None:
    """XLSX/CSV attachment → model-readable row context with sheet/row provenance.

    Preserves sheet and row provenance; does not infer or calculate domain values.
    """
    suffix = path.suffix.lower()
    parts: list[str] = [
        f"Файл: {original_name}",
        "Тип данных: таблица/спецификация. Строки ниже сохранены с номерами листов/строк.",
        "",
    ]
    truncated = False
    used_chars = len('\n'.join(parts))

    def _append(line: str) -> bool:
        nonlocal truncated, used_chars
        extra_chars = len(line) + (1 if used_chars else 0)
        if used_chars + extra_chars > max_chars:
            truncated = True
            return False
        parts.append(line)
        used_chars += extra_chars
        return True

    if suffix in {".xlsx", ".xlsm"}:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        except Exception as err:  # noqa: BLE001
            logger.warning("[ATTACH] xlsx structured read failed for %s: %s", original_name, err)
            return None
        try:
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                if not _append(f"## Лист: {sheet}"):
                    break
                nonempty = 0
                for ri, row in enumerate(ws.iter_rows(values_only=True), 1):
                    vals = [_compact_cell(v) for v in row if v not in (None, "")]
                    if not vals:
                        continue
                    nonempty += 1
                    if not _append(f"{sheet}!R{ri}: " + " | ".join(vals)):
                        break
                _append(f"Итого непустых строк на листе «{sheet}»: {nonempty}")
                if truncated:
                    break
        finally:
            wb.close()
    elif suffix == ".csv":
        import csv
        import io
        from backend.text_decoding import read_document_text
        text = read_document_text(path)
        sample = text.partition('\n')[0]
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        nonempty = 0
        _append("## CSV")
        for ri, row in enumerate(reader, 1):
            vals = [_compact_cell(v) for v in row if str(v or "").strip()]
            if not vals:
                continue
            nonempty += 1
            if not _append(f"CSV!R{ri}: " + " | ".join(vals)):
                break
        _append(f"Итого непустых строк CSV: {nonempty}")
    else:
        return None

    if truncated and parts[-1] != "[Табличный контекст усечён по лимиту; для полного документа нужен датасет или более узкий лист.]":
        parts.append("[Табличный контекст усечён по лимиту; для полного документа нужен датасет или более узкий лист.]")
    text = "\n".join(parts).strip()
    return (text, truncated) if text else None

async def _prepare_read_attachment(
    temp_path: Path,
    original_name: str,
    *,
    attachment_id: str | None = None,
) -> dict[str, Any]:
    """Convert and preserve one lossless, server-owned read attachment."""
    from uuid import uuid4

    attach_id = attachment_id or f"read_{uuid4().hex[:12]}"
    from backend.product_edition import is_light
    from proxy.services.chat_attachment_service import IMAGE_SUFFIXES, image_data_url, preserve_read_attachment
    if is_light() and Path(original_name).suffix.lower() in IMAGE_SUFFIXES:
        try:
            await asyncio.to_thread(image_data_url, temp_path)
        except (OSError, ValueError) as error:
            raise HTTPException(422, f"Не удалось прочитать изображение: {error}") from error
        await asyncio.to_thread(preserve_read_attachment, temp_path, attachment_id=attach_id, original_name=original_name)
        return {"attachment_id": attach_id, "mode": "read", "name": original_name,
                "chars": 0, "text": "", "truncated": False, "media_type": "image"}
    try:
        structured = await asyncio.to_thread(
            _format_tabular_attachment_context,
            temp_path,
            original_name,
            max_chars=_READ_ATTACH_MAX_CHARS,
        )
    except (OSError, ValueError) as error:
        raise HTTPException(422, f'Не удалось прочитать таблицу «{original_name}»: {error}') from error
    if structured:
        text, truncated = structured
    else:
        from backend.converter import convert_to_markdown

        try:
            text = await asyncio.to_thread(convert_to_markdown, temp_path)
        except Exception as error:  # noqa: BLE001 - upload errors must stay operator-visible
            logger.warning("[ATTACH] read conversion failed for %s: %s", original_name, error)
            raise HTTPException(
                422,
                f"Не удалось прочитать файл «{original_name}»: {error}. "
                "Попробуй режим индексации/OCR или другой формат файла.",
            ) from error
        text = (text or "").strip()
        truncated = len(text) > _READ_ATTACH_MAX_CHARS

    text = (text or "").strip()
    suffix = Path(original_name).suffix.lower()
    if not text:
        raise HTTPException(
                422,
                f"Не удалось прочитать текст из «{original_name}». "
                "Для таблиц попробуй режим быстрой сверки, для сканов — индексацию/OCR.",
        )

    from proxy.services.chat_attachment_service import cleanup_expired, preserve_read_attachment

    await asyncio.to_thread(cleanup_expired)
    await asyncio.to_thread(
        preserve_read_attachment,
        temp_path,
        attachment_id=attach_id,
        original_name=original_name,
    )
    return {
        "attachment_id": attach_id,
        "mode": "read",
        "name": original_name,
        "chars": len(text),
        "text": text[:_READ_ATTACH_MAX_CHARS],
        "truncated": truncated,
    }
