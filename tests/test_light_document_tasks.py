import asyncio
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
import json
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from proxy.services.document_task_store import DocumentTaskStore
from proxy.services.tabular_document_service import read_table, preview, encode
from proxy.services.table_document_tool import bind_table_tool, without_table_replays, TableDocumentTool


def workbook(path, count=70):
    book = Workbook()
    sheet = book.active
    sheet.title = "Лист '1'"
    sheet.append(['Название', 'Количество', 'Единица', 'Примечание'])
    for index in range(count):
        sheet.append([f'Позиция {index}', index, 'м', 'Всего в упаковке 100 м'])
    book.save(path)
    book.close()
    return read_table(path)


def decisions(packet, **change):
    return [{'row_id': row['id'], 'status': 'reviewed', 'note': 'Прочитано без изменения исходника',
             'cells': [row['cells'][0]['coordinate']], **change} for row in packet['rows']]


def test_literal_cells_zeros_formula_and_multiple_sheets(tmp_path):
    path = tmp_path/'Книга с пробелами.xlsx'
    book = Workbook()
    sheet = book.active
    sheet.append(['Наименование', None, 'Кол-во', 'Ед.'])
    sheet.append(['Монтаж', None, 0, None])
    sheet.append(['Всего в упаковке', False, '=1+1', '#DIV/0!'])
    sheet.merge_cells('A5:C5')
    sheet['A5'] = 'Заголовок'
    extra = book.create_sheet('Другой лист')
    extra.sheet_state = 'hidden'
    extra.append(['Повтор', 0.001])
    book.save(path)
    book.close()
    doc = read_table(path)
    assert doc['row_count'] == 5
    assert [c['coordinate'] for c in doc['rows'][1]['cells']] == ['A2', 'C2']
    assert doc['rows'][1]['cells'][1]['value'] == '0'
    assert doc['rows'][2]['cells'][1]['value'] is False
    assert doc['rows'][2]['cells'][2] == {
        'coordinate': 'C3', 'kind': 'formula', 'formula': '=1+1', 'cached': None, 'needs_recalculation': True}
    assert doc['rows'][2]['cells'][3]['kind'] == 'error'
    assert doc['sheets'][0]['merged_ranges'] == ['A5:C5']
    assert doc['sheets'][1]['hidden'] is True
    assert doc['rows'][-1]['id'] == 's2:r1'


@pytest.mark.parametrize('encoding', ['utf-8-sig', 'utf-16'])
def test_csv_coordinates_and_multiline_are_literal(tmp_path, encoding):
    path = tmp_path/'таблица.csv'
    path.write_text('Имя;Пусто;Число\n"две\nстроки";;0\n', encoding=encoding, newline='')
    doc = read_table(path)
    assert doc['rows'][1]['cells'] == [
        {'coordinate': 'A2', 'kind': 'text', 'value': 'две\nстроки'},
        {'coordinate': 'C2', 'kind': 'text', 'value': '0'}]


def test_preview_never_clips_cell_or_claims_full_extraction(tmp_path):
    path = tmp_path/'long.csv'
    path.write_text('Title\n' + 'д' * 9000, encoding='utf-8')
    doc = read_table(path)
    text, truncated = preview(doc, 1200)
    assert truncated and len(text) <= 1200
    assert 'д' * 100 not in text
    assert len(doc['rows'][1]['cells'][0]['value']) == 9000


def test_resume_300_rows_is_exact_and_idempotent(tmp_path):
    doc = workbook(tmp_path/'source.xlsx', 300)
    db = tmp_path/'tasks.db'
    store = DocumentTaskStore(db)
    task = store.open('scope', 'Проверить все строки', doc)
    original = encode(doc)
    previous, seen = None, []
    while store.status('scope', task)['pending']:
        packet = store.next_packet('scope', task, max_bytes=3000)
        assert len(encode(packet['rows']).encode('utf-8')) <= 3000
        # An interrupted model request has not advanced coverage.
        store = DocumentTaskStore(db)
        assert store.next_packet('scope', task, max_bytes=3000) == packet
        data = decisions(packet)
        store.commit('scope', task, packet['packet_id'], data)
        store.commit('scope', task, packet['packet_id'], data)
        if previous:
            store.commit('scope', task, previous[0], previous[1])
        previous = packet['packet_id'], data
        seen += [row['id'] for row in packet['rows']]
    status = store.status('scope', task)
    assert status['coverage_complete'] and status['reviewed'] == 301
    assert len(set(seen)) == 301
    assert store.open('scope', 'Проверить все строки', doc) == task
    assert encode(doc) == original
    # Closed connections must not keep this file locked on Windows.
    db.rename(tmp_path/'moved.db')


def test_atomic_commit_rejects_omissions_duplicates_and_fabricated_cells(tmp_path):
    store = DocumentTaskStore(tmp_path/'tasks.db')
    task = store.open('scope', 'Check', workbook(tmp_path/'source.xlsx', 4))
    packet = store.next_packet('scope', task)
    good = decisions(packet)
    bad = [*good[:-1], {**good[-1], 'cells': ['Z999']}]
    for rows in (good[:-1], [good[0]] * len(good), bad):
        with pytest.raises(ValueError):
            store.commit('scope', task, packet['packet_id'], rows)
        assert store.status('scope', task)['pending'] == 5
    with pytest.raises(ValueError):
        store.commit('scope', task, packet['packet_id'], [{**row, 'quantity': 9} for row in good])
    store.commit('scope', task, packet['packet_id'], good)
    with pytest.raises(ValueError, match='другим'):
        store.commit('scope', task, packet['packet_id'], decisions(packet, note='Changed'))


def test_scope_and_source_changes_never_inherit_results(tmp_path, monkeypatch):
    monkeypatch.setenv('LES_PRODUCT_EDITION', 'light')
    monkeypatch.setenv('LES_WINDOWS_STATE_ROOT', str(tmp_path/'state'))
    from proxy.services.chat_attachment_service import preserve_read_attachment
    path = tmp_path/'source.xlsx'
    workbook(path, 2)
    preserve_read_attachment(path, attachment_id='read_aaaaaaaaaaaa', original_name=path.name)
    req = SimpleNamespace(attachment_id='read_aaaaaaaaaaaa', session_id='session', question='Check',
                          project_id=1, scope={'scope_type': 'datasets'})
    def bind(**changes):
        return bind_table_tool(req, **({'dataset_ids': ['A'], 'profile_revision': 'p1',
                                       'model_revision': 'm1', 'input_budget': 24000} | changes))
    a = bind()
    packet = a.execute({'operation': 'next'})
    a.execute({'operation': 'record', 'packet_id': packet['packet_id'], 'decisions': decisions(packet)})
    assert bind().state()['coverage_complete']
    for changes in ({'dataset_ids': ['B']}, {'profile_revision': 'p2'}, {'model_revision': 'm2'}):
        b = bind(**changes)
        assert b.state()['pending'] == 3
        with pytest.raises(ValueError):
            b.execute({'operation': 'results', 'task_id': a.task_id})
    req.session_id = 'other'
    assert bind().state()['pending'] == 3
    req.session_id = 'session'
    req.question = 'Other question'
    b = bind()
    new_task_id = b.task_id
    assert b.state()['pending'] == 3
    assert b.execute({'operation': 'status', 'task_id': a.task_id})['coverage_complete']
    req.attachment_id = None
    assert bind().task_id == a.task_id  # Explicitly resumed task remains active after restart.
    restarted = bind()
    assert restarted.execute({'operation': 'start'})['task_id'] == new_task_id
    assert restarted.state()['pending'] == 3
    assert bind(dataset_ids=['never-selected']) is None


def test_uncomputed_formula_remains_a_visible_issue(tmp_path):
    path = tmp_path/'formula.xlsx'
    book = Workbook(); book.active.append(['Title', '=1+1']); book.save(path); book.close()
    store = DocumentTaskStore(tmp_path/'tasks.db')
    task = store.open('scope', 'Check', read_table(path))
    packet = store.next_packet('scope', task)
    with pytest.raises(ValueError, match='needs_review'):
        store.commit('scope', task, packet['packet_id'], decisions(packet))
    store.commit('scope', task, packet['packet_id'], decisions(packet, status='needs_review'))
    assert store.status('scope', task)['needs_attention']


def test_oversized_row_is_blocked_not_truncated_or_skipped(tmp_path):
    path = tmp_path/'long.csv'; path.write_text('x'*12000, encoding='utf-8')
    doc = read_table(path)
    store = DocumentTaskStore(tmp_path/'tasks.db'); task = store.open('scope', 'Read', doc)
    packet = store.next_packet('scope', task, max_bytes=512)
    assert packet['blocked_row'] == 's1:r1' and packet['packet_id'] is None
    assert store.status('scope', task)['pending'] == 1


def test_corrupt_workbook_is_rejected_as_read_error(tmp_path):
    from fastapi import HTTPException
    from proxy.services.chat_attachment_read_service import _prepare_read_attachment
    path = tmp_path/'broken.xlsx'; path.write_bytes(b'not a workbook')
    with pytest.raises(HTTPException) as caught:
        asyncio.run(_prepare_read_attachment(path, path.name))
    assert caught.value.status_code == 422


def test_cached_formula_error_cannot_be_marked_reviewed(tmp_path):
    from io import BytesIO
    from zipfile import ZipFile
    path = tmp_path/'error.xlsx'
    book = Workbook(); book.active['A1'] = '=1/0'; book.save(path); book.close()
    original = path.read_bytes()
    with ZipFile(BytesIO(original)) as source, ZipFile(path, 'w') as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == 'xl/worksheets/sheet1.xml':
                data = data.replace(b'<c r="A1">', b'<c r="A1" t="e">').replace(b'<v></v>', b'<v>#DIV/0!</v>')
            target.writestr(item, data)
    doc = read_table(path)
    assert doc['rows'][0]['cells'][0]['cached']['kind'] == 'error'
    store = DocumentTaskStore(tmp_path/'tasks.db'); task = store.open('scope', 'Check', doc)
    packet = store.next_packet('scope', task)
    with pytest.raises(ValueError, match='needs_review'):
        store.commit('scope', task, packet['packet_id'], decisions(packet))


def test_concurrent_redelivery_does_not_double_count(tmp_path):
    doc = workbook(tmp_path/'source.xlsx', 1)
    store = DocumentTaskStore(tmp_path/'tasks.db'); task = store.open('scope', 'Check', doc)
    packet = store.next_packet('scope', task)
    def commit(_):
        store.commit('scope', task, packet['packet_id'], decisions(packet))
    with ThreadPoolExecutor(max_workers=3) as pool:
        list(pool.map(commit, range(6)))
    assert store.status('scope', task)['revision'] == 1


def test_bound_tool_executes_through_real_harness_and_context_stays_bounded(tmp_path, monkeypatch):
    monkeypatch.setenv('LES_PRODUCT_EDITION', 'light')
    monkeypatch.setenv('LES_WINDOWS_STATE_ROOT', str(tmp_path/'state'))
    from proxy.services.tool_harness_service import ToolHarness
    doc = workbook(tmp_path/'source.xlsx', 70)
    store = DocumentTaskStore(tmp_path/'tasks.db'); task = store.open('scope', 'Check', doc)
    tool = TableDocumentTool(store, 'scope', task, doc, 4000)
    harness = ToolHarness(); tool.register(harness)
    logs = []
    while tool.state()['pending']:
        payload = asyncio.run(harness.call_async('table_document', {'operation': 'next'}))
        assert payload['status'] == 'ok', payload
        packet = payload['result']; logs.append(payload)
        saved = asyncio.run(harness.call_async('table_document', {
            'operation': 'record', 'packet_id': packet['packet_id'], 'decisions': decisions(packet)}))
        assert saved['status'] == 'ok', saved
        logs.append(saved)
        assert len(encode(tool.context())) < 2000
    assert without_table_replays(logs) == []
    assert tool.state()['coverage_complete']


def test_current_packet_is_mandatory_and_never_sliced():
    from proxy.services.chat_evidence_context import govern_inference_messages
    from proxy.services.model_execution_preset_service import _FACTORY_9B
    from proxy.services.context_governor_service import ContextRequiredSectionOverflow
    preset = replace(_FACTORY_9B, input_token_limit=500)
    with pytest.raises(ContextRequiredSectionOverflow):
        govern_inference_messages(preset=preset, profile_prefix='Read', request_payload='Check',
                                  required_evidence=[{'cell': 'Ж' * 10000}])


def test_model_sees_started_task_and_next_action_without_replayed_history(tmp_path):
    doc = workbook(tmp_path/'table.xlsx', 1)
    store = DocumentTaskStore(tmp_path/'tasks.db'); task = store.open('scope', 'Check', doc)
    tool = TableDocumentTool(store, 'scope', task, doc, 4000, question='Check')
    assert '0 из 2' in tool.coverage_notice()
    assert tool.context()['task_is_for_current_question'] is True
    assert tool.context()['next_action_for_active_task'] == 'next'
    with pytest.raises(ValueError, match='already started'):
        tool.execute({'operation': 'start'})
    assert tool.context()['next_action_for_active_task'] == 'next'
    packet = tool.execute({'operation': 'next'})
    with pytest.raises(ValueError, match='Use record'):
        tool.execute({'operation': 'start'})
    assert tool.context()['current_read']['packet_id'] == packet['packet_id']
    assert tool.context()['next_action_for_active_task'] == 'record'
    tool.execute({'operation': 'record', 'packet_id': packet['packet_id'], 'decisions': decisions(packet)})
    assert tool.context()['next_action_for_active_task'] == 'finish_or_read_results'
    assert tool.coverage_notice() == ''
    tool.question = 'Different objective'
    assert tool.context()['task_is_for_current_question'] is False


def test_native_tool_turn_is_bounded_and_record_uses_server_owned_packet(tmp_path, monkeypatch):
    from proxy.services.tool_harness_service import ToolHarness
    from proxy.services.chat_evidence_context import govern_inference_messages
    from proxy.services.model_execution_preset_service import _FACTORY_9B
    from proxy.services.context_governor_service import ContextRequiredSectionOverflow
    monkeypatch.setenv('LES_PRODUCT_EDITION', 'light')
    doc = workbook(tmp_path/'source.xlsx', 1)
    store = DocumentTaskStore(tmp_path/'tasks.db'); task = store.open('scope', 'Check', doc)
    tool = TableDocumentTool(store, 'scope', task, doc, 4000, question='Check')
    harness = ToolHarness(); tool.register(harness)
    packet = asyncio.run(harness.call_async('table_document', {'operation': 'next'}))['result']
    messages, _ = govern_inference_messages(preset=_FACTORY_9B, profile_prefix='Read',
                                            request_payload='Check', **tool.model_context())
    assert messages[-2]['role'] == 'assistant' and messages[-1]['role'] == 'tool'
    assert messages[-2]['tool_calls'][0]['id'] == messages[-1]['tool_call_id']
    assert packet['packet_id'] in messages[-1]['content']
    wrong = [{**row, 'cells': ['cell content']} for row in decisions(packet)]
    rejected = asyncio.run(harness.call_async('table_document', {'operation': 'record', 'decisions': wrong}))
    assert rejected['status'] == 'error'
    tool.remember_validation_failure({'operation': 'record', 'decisions': wrong}, rejected)
    assert json.loads(tool.last_exchange[-1]['content'])['status'] == 'error'
    assert tool.state()['pending'] == 2
    with pytest.raises(ContextRequiredSectionOverflow):
        govern_inference_messages(preset=replace(_FACTORY_9B, input_token_limit=100),
                                  profile_prefix='Read', request_payload='Check', **tool.model_context())
    failed = asyncio.run(harness.call_async('table_document', {'operation': 'record', 'packet_id': 'fake',
                                'decisions': decisions(packet)}))
    assert failed['status'] == 'error' and tool.state()['pending'] == 2
    assert json.loads(tool.last_exchange[-1]['content'])['status'] == 'error'
    saved = asyncio.run(harness.call_async('table_document', {'operation': 'record', 'decisions': decisions(packet)}))
    assert saved['status'] == 'ok' and tool.state()['coverage_complete']
    assert len(tool.last_exchange) == 2 and 'fake' not in tool.last_exchange[-1]['content']
